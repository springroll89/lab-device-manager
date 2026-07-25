import json
import time
from io import BytesIO

from openpyxl import load_workbook
from lab_device_manager.web.app import create_app
from lab_device_manager.runtime.types import DeviceConfig
from lab_device_manager.db.repository import Repository
from lab_device_manager.instruments.base import StatusSnapshot


def _sample_to_dict(s):
    return {"id": s.id, "run_id": s.run_id, "device_id": s.device_id,
            "ts_ms": s.ts_ms, "state": s.state, "flow_rate": s.flow_rate,
            "delivered_volume": s.delivered_volume, "temp_c": s.temp_c,
            "metrics": json.loads(s.metrics_json) if s.metrics_json else {}}


def _event_to_dict(e):
    return {"id": e.id, "device_id": e.device_id, "run_id": e.run_id,
            "ts_ms": e.ts_ms, "event_type": e.event_type, "severity": e.severity,
            "detail": json.loads(e.detail_json) if e.detail_json else {}}


class StaticEngine:
    def __init__(self, latest, dmap):
        self._latest = latest; self._dmap = dmap
    def latest(self): return dict(self._latest)
    def device_map(self): return dict(self._dmap)


def _snap(state):
    return StatusSnapshot(timestamp=time.time(), state=state, work_mode="仅注入",
                          device_id="d", acc_volume=12.5, acc_unit="mL", temp_c=40.0,
                          metrics={"syringe_code": 9, "target_volume": 50.0,
                                   "target_unit": "mL", "inject_rate": 5.0,
                                   "pause_delay_ms": 5000, "repeat_count": 3, "force": 100})


def _app(tmp_path, *, auth_bypass: bool = True):
    repo = Repository(":memory:")
    did = repo.upsert_device("pump-1", "tyd02", alias="注射泵")
    eng = StaticEngine({did: _snap("running")}, {did: DeviceConfig(name="pump-1", type="tyd02", alias="注射泵")})
    app = create_app(eng, repo, secret_key="test-secret")
    app.config.update(TESTING=True, AUTH_TEST_BYPASS=auth_bypass)
    return app, repo, did


def test_status_lists_devices(tmp_path):
    app, repo, did = _app(tmp_path)
    r = app.test_client().get("/api/status")
    data = r.get_json()
    assert data["devices"][0]["name"] == "pump-1"
    assert data["devices"][0]["latest"]["state"] == "running"


def test_runs_and_tag(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    untagged = app.test_client().get("/api/runs/untagged").get_json()
    assert len(untagged) == 1 and untagged[0]["tagged"] is False
    app.test_client().post(f"/api/runs/{rid}/tag", json={"operator": "alice", "project_tag": "陶瓷膜-A"})
    assert app.test_client().get("/api/runs/untagged").get_json() == []


def test_tag_missing_run_returns_404(tmp_path):
    app, repo, did = _app(tmp_path)
    response = app.test_client().post(
        "/api/runs/99999/tag",
        json={"operator": "alice"},
    )
    assert response.status_code == 404


def test_runs_csv_export(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    r = app.test_client().get("/api/runs/export.csv")
    assert r.status_code == 200
    assert "text/csv" in r.content_type
    assert r.headers.get("Content-Disposition") == "attachment; filename=pump_runs.csv"
    body = r.data.decode("utf-8")
    assert "end_status" in body and "completed" in body


def test_runs_include_setpoints_and_actual(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {"syringe_code": 9, "target_volume": 50.0})
    repo.close_run(rid, 1751000060_000, "completed", 50.0, "mL", 150.0, "mL", 0)
    r = app.test_client().get("/api/runs").get_json()[0]
    assert r["actual_volume"] == 50.0
    assert r["setpoints"]["syringe_code"] == 9


def test_csv_export_neutralizes_formula_injection(tmp_path):
    """CSV injection (OWASP): user-entered fields starting with = + - @ must be
    prefixed so spreadsheet apps don't evaluate them as formulas."""
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    repo.tag_run(rid, "=cmd|'/c calc'!A1", "+2+2", "-3", "@sum(A1)")
    body = app.test_client().get("/api/runs/export.csv").data.decode("utf-8")
    assert "'=cmd" in body      # = neutralized
    assert "'+2+2" in body      # + neutralized
    assert "'-3" in body        # - neutralized
    assert "'@sum" in body      # @ neutralized


def test_device_page_serves_html(tmp_path):
    app, repo, did = _app(tmp_path)
    r = app.test_client().get("/device/" + str(did))
    assert r.status_code == 200
    assert b"device.js" in r.data


def test_device_detail_includes_metrics_in_latest(tmp_path):
    """_snap_to_dict must expose metrics so the dashboard config line + detail
    page can render mode/syringe/target without a second round-trip."""
    app, repo, did = _app(tmp_path)
    data = app.test_client().get("/api/devices/" + str(did)).get_json()
    assert data["device"]["name"] == "pump-1"
    assert data["latest"]["metrics"]["syringe_code"] == 9
    assert data["latest"]["metrics"]["target_volume"] == 50.0
    assert data["metrics"]["syringe_code"] == 9


def test_whd_realtime_data_reads_latest_snapshot_without_polling_or_writing():
    repo = Repository(":memory:")
    device_id = repo.upsert_device("whd-1", "whd46", alias="环境温湿度")
    snapshot = StatusSnapshot(
        timestamp=time.time(),
        state="running",
        work_mode="monitoring",
        device_id="whd",
        temp_c=25.0,
        metrics={
            "channels": [
                {"temp": 24.0, "humid": 40.0},
                {"temp": 25.0, "humid": 41.0},
                {"temp": 26.0, "humid": 42.0},
            ],
            "avg_humid_rh": 41.0,
        },
    )

    class PollingMustNotRun:
        is_connected = True

        def read_channels(self):
            raise AssertionError("GET realtime-data must not poll the serial adapter")

    class WHDEngine(StaticEngine):
        def _get_adapter(self, requested_id):
            assert requested_id == device_id
            return PollingMustNotRun()

    engine = WHDEngine(
        {device_id: snapshot},
        {
            device_id: DeviceConfig(
                name="whd-1", type="whd46", alias="环境温湿度"
            )
        },
    )
    app = create_app(engine, repo, secret_key="test-secret")
    app.config.update(TESTING=True, AUTH_TEST_BYPASS=True)

    response = app.test_client().get(
        f"/api/devices/{device_id}/realtime-data"
    )

    assert response.status_code == 200
    body = response.get_json()
    assert body["channels"][0] == {
        "channel": 1,
        "temp_c": 24.0,
        "humid_rh": 40.0,
    }
    assert repo.list_samples_for_device(device_id) == []


def test_whd_connect_hands_verified_port_to_managed_engine(monkeypatch):
    repo = Repository(":memory:")
    device_id = repo.upsert_device(
        "whd-1",
        "whd46",
        alias="环境温湿度",
    )

    class ManagedEngine(StaticEngine):
        def __init__(self):
            super().__init__(
                {},
                {
                    device_id: DeviceConfig(
                        name="whd-1",
                        type="whd46",
                        alias="环境温湿度",
                        serial_port="/dev/old",
                        parity="NONE",
                    )
                },
            )
            self.reconnected = None

        def reconnect_device(self, requested_id, port):
            self.reconnected = (requested_id, port)
            return {"ok": True, "port": port}

    engine = ManagedEngine()
    monkeypatch.setattr(
        "lab_device_manager.web.whd46.probe_port",
        lambda *args, **kwargs: {
            "ok": True,
            "port": "/dev/new",
            "error": "",
            "detail": "",
        },
    )
    app = create_app(engine, repo, secret_key="test-secret")
    app.config.update(TESTING=True, AUTH_TEST_BYPASS=True)

    response = app.test_client().post(
        f"/api/devices/{device_id}/connect",
        json={"port": "/dev/new"},
    )

    assert response.status_code == 200
    assert response.get_json()["port"] == "/dev/new"
    assert engine.reconnected == (device_id, "/dev/new")


def test_device_detail_filters_runs_by_device_id(tmp_path):
    """SECURITY: /api/devices/<id> must only return that device's runs — never
    another device's. (Regression: previously returned repo.list_runs for ALL
    devices.)"""
    app, repo, did = _app(tmp_path)
    other = repo.upsert_device("pump-2", "tyd02")
    r1 = repo.open_run(did, 1, 1751000000_000, {})
    r2 = repo.open_run(other, 1, 1751000100_000, {})
    repo.close_run(r1, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    repo.close_run(r2, 1751000160_000, "completed", 99.0, "mL", 99.0, "mL", 0)
    data = app.test_client().get("/api/devices/" + str(did)).get_json()
    run_ids = [r["id"] for r in data["runs"]]
    assert run_ids == [r1]              # only this device's run
    assert all(r["device_id"] == did for r in data["runs"])


def test_runs_filter_by_project_tag(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    repo.tag_run(rid, "alice", "陶瓷膜-A", "", "")
    data = app.test_client().get("/api/runs?project_tag=陶瓷膜-A").get_json()
    assert len(data) == 1 and data[0]["project_tag"] == "陶瓷膜-A"
    data = app.test_client().get("/api/runs?project_tag=不存在").get_json()
    assert data == []


def test_run_detail_includes_samples_and_events(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.add_sample(rid, did, 1751000001_000, "running", 1.2, 1.0, 33.0, "{}")
    repo.add_event(did, rid, 1751000000_000, "start", "info", "{}")
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    data = app.test_client().get(f"/api/runs/{rid}").get_json()
    assert data["run"]["id"] == rid
    assert len(data["samples"]) == 1
    assert data["samples"][0]["flow_rate"] == 1.2
    assert len(data["events"]) == 1
    assert data["events"][0]["event_type"] == "start"


def test_run_detail_404_for_missing_run(tmp_path):
    app, repo, did = _app(tmp_path)
    r = app.test_client().get("/api/runs/99999")
    assert r.status_code == 404


def test_runs_export_xlsx(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000,
                       {"syringe_code": 9, "target_volume": 50.0})
    repo.add_sample(rid, did, 1751000001_000, "running", 1.2, 1.0, 33.0, "{}")
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    r = app.test_client().get("/api/runs/export.xlsx")
    assert r.status_code == 200
    assert r.headers.get("Content-Disposition") == "attachment; filename=pump_runs.xlsx"
    wb = load_workbook(BytesIO(r.data))
    assert "runs" in wb.sheetnames and "samples" in wb.sheetnames
    ws = wb["runs"]
    assert ws.cell(row=1, column=1).value == "id"
    assert any(ws.cell(row=row, column=1).value == rid for row in range(2, ws.max_row + 1))


def test_runs_export_xlsx_rejects_invalid_numeric_filters(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()
    assert client.get("/api/runs/export.xlsx?limit=abc").status_code == 400
    assert client.get("/api/runs/export.xlsx?offset=-1").status_code == 400
    assert client.get("/api/runs/export.xlsx?device_id=abc").status_code == 400


def test_directory_browser_rejects_paths_outside_storage_roots(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    app, repo, did = _app(tmp_path)
    response = app.test_client().get(
        "/api/browse-directory?path=/etc"
    )
    assert response.status_code == 400


def test_directory_browser_lists_workspace_root(tmp_path, monkeypatch):
    (tmp_path / "exports").mkdir()
    monkeypatch.chdir(tmp_path)
    app, repo, did = _app(tmp_path)
    response = app.test_client().get(
        f"/api/browse-directory?path={tmp_path}"
    )
    assert response.status_code == 200
    assert "exports" in response.get_json()["directories"]


def test_run_samples_endpoint(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.add_sample(rid, did, 1751000001_000, "running", 1.2, 1.0, 33.0, "{}")
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    data = app.test_client().get(f"/api/runs/{rid}/samples").get_json()
    assert len(data) == 1
    assert data[0]["flow_rate"] == 1.2
    assert app.test_client().get("/api/runs/99999/samples").status_code == 404


def test_run_events_endpoint(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    repo.add_event(did, rid, 1751000000_000, "start", "info", "{}")
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    data = app.test_client().get(f"/api/runs/{rid}/events").get_json()
    assert len(data) == 1
    assert data[0]["event_type"] == "start"
    assert app.test_client().get("/api/runs/99999/events").status_code == 404


def test_runs_query_params_validated(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()
    assert client.get("/api/runs?limit=abc").status_code == 400
    assert client.get("/api/runs?offset=-1").status_code == 400
    assert client.get("/api/runs?device_id=abc").status_code == 400
    assert client.get("/api/runs?start_ms=abc").status_code == 400
    assert client.get("/api/runs?limit=1001").status_code == 400
    # valid params still work
    assert client.get("/api/runs?limit=10&offset=0").status_code == 200


def test_run_report_pdf(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {"target_volume": 50.0})
    repo.add_sample(rid, did, 1751000001_000, "running", 1.2, 1.0, 33.0, "{}")
    repo.add_event(did, rid, 1751000000_000, "start", "info", "{}")
    repo.close_run(rid, 1751000060_000, "completed", 12.5, "mL", 12.5, "mL", 0)
    r = app.test_client().get(f"/api/runs/{rid}/report.pdf")
    assert r.status_code == 200
    assert r.content_type == "application/pdf"
    assert r.data[:4] == b"%PDF"


def test_run_page_serves_html(tmp_path):
    app, repo, did = _app(tmp_path)
    rid = repo.open_run(did, 1, 1751000000_000, {})
    r = app.test_client().get(f"/run/{rid}")
    assert r.status_code == 200
    assert b"run.js" in r.data


def test_shared_premium_theme_is_served_by_every_operator_page(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()

    theme = client.get("/static/puricore-theme.css")
    assert theme.status_code == 200
    assert b"--pc-accent" in theme.data
    assert b"prefers-reduced-motion" in theme.data

    for page in (
        "index.html",
        "device.html",
        "run.html",
        "sensor.html",
        "login.html",
        "experiment.html",
        "accounts.html",
        "change-password.html",
    ):
        response = client.get(f"/static/{page}")
        assert response.status_code == 200
        assert b"/static/puricore-theme.css" in response.data


def test_dashboard_uses_the_same_outer_card_style_for_every_device(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()

    page = client.get("/static/index.html")
    script = client.get("/static/app.js")
    theme = client.get("/static/puricore-theme.css")

    assert b".card.sensor-card" not in page.data
    assert b"sensor-card" not in script.data
    assert b".page-dashboard .card.sensor-card" not in theme.data


def test_account_creation_keeps_stable_form_reference(tmp_path):
    app, repo, did = _app(tmp_path)
    script = app.test_client().get("/static/accounts.js")

    assert script.status_code == 200
    assert b"const form = event.currentTarget" in script.data
    assert b"form.reset()" in script.data
    assert b"event.currentTarget.reset()" not in script.data


def test_authenticated_pages_share_account_menu_entry(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()

    menu_script = client.get("/static/account-menu.js")
    assert menu_script.status_code == 200
    assert b'addMenuLink(panel, "/change-password", "' in menu_script.data
    assert b'addMenuLink(panel, "/experiments", "' not in menu_script.data
    assert b'addMenuLink(panel, "/inventory", "' not in menu_script.data
    assert b"session.can_manage_accounts" in menu_script.data
    assert b'addMenuLink(panel, "/accounts", "' in menu_script.data

    for page in (
        "index.html",
        "device.html",
        "run.html",
        "sensor.html",
        "experiment.html",
        "materials.html",
        "hazardous-waste.html",
        "measurement-station.html",
        "accounts.html",
    ):
        response = client.get(f"/static/{page}")
        assert response.status_code == 200
        assert b"data-account-menu" in response.data
        assert b"/static/account-menu.js" in response.data


def test_authenticated_pages_share_fixed_primary_navigation(tmp_path):
    app, repo, did = _app(tmp_path)
    client = app.test_client()

    navigation = client.get("/static/primary-navigation.js")
    assert navigation.status_code == 200
    labels = (
        "设备看板",
        "实验执行",
        "物品与库存",
        "危废管理",
        "粘度工位",
    )
    positions = [navigation.data.index(label.encode()) for label in labels]
    assert positions == sorted(positions)
    assert b'link.setAttribute("aria-current", "page")' in navigation.data

    for page in (
        "index.html",
        "device.html",
        "run.html",
        "sensor.html",
        "experiment.html",
        "materials.html",
        "hazardous-waste.html",
        "measurement-station.html",
        "accounts.html",
    ):
        response = client.get(f"/static/{page}")
        assert response.status_code == 200
        assert b"data-primary-nav" in response.data
        assert b"/static/primary-navigation.js" in response.data

    experiment = client.get("/static/experiment.html").data
    header = experiment.split(b"</header>", 1)[0]
    assert b"cameraScanButton" not in header
    assert "纸电并行验证".encode() not in header
    assert b"cameraScanButton" in experiment
    assert "纸电并行验证".encode() in experiment


def _authenticated_app(tmp_path):
    return _app(tmp_path, auth_bypass=False)


def _login(client, username="admin", password="admin"):
    return client.post(
        "/login",
        json={"username": username, "password": password},
    )


def _change_password(client, current_password, new_password):
    session_data = client.get("/api/session").get_json()
    return client.post(
        "/api/account/password",
        json={
            "current_password": current_password,
            "new_password": new_password,
        },
        headers={"X-CSRF-Token": session_data["csrf_token"]},
    )


def _experiment_request(batch_id="20260725-CEM-01"):
    return {
        "batch_id": batch_id,
        "membrane_system": "CEM",
        "recipe_no": "R-CEM-001",
        "recipe_version": "V1",
        "sop_code": "SOP-SOL-GEL-CEM-AEM-01",
        "sop_version": "V0.2",
        "target_viscosity_min_mpas": 2.5,
        "target_viscosity_max_mpas": 10,
        "reviewer": "",
        "spec_snapshot": {},
    }


def test_login_is_always_required_for_pages_and_apis(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()

    page = client.get("/")
    device = client.get(f"/device/{did}")
    api = client.get("/api/status")
    direct_static_page = client.get("/static/index.html")
    theme = client.get("/static/puricore-theme.css")

    assert page.status_code == 302
    assert page.headers["Location"].startswith("/login?next=")
    assert device.status_code == 302
    assert api.status_code == 401
    assert api.get_json()["error"] == "unauthorized"
    assert direct_static_page.status_code == 302
    assert theme.status_code == 200


def test_login_page_and_account_auth_mode_are_public(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()

    page = client.get("/login")
    mode = client.get("/api/auth-mode")

    assert page.status_code == 200
    assert "账号".encode() in page.data
    assert "密码".encode() in page.data
    assert mode.get_json() == {
        "password_required": True,
        "account_auth": True,
        "initial_username": "admin",
    }


def test_initial_admin_login_requires_immediate_password_change(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()

    login = _login(client)
    assert login.status_code == 200
    assert login.get_json()["role"] == "super_admin"
    assert login.get_json()["must_change_password"] is True

    page = client.get("/")
    api = client.get("/api/status")
    password_page = client.get("/change-password")

    assert page.status_code == 302
    assert page.headers["Location"].startswith("/change-password?next=")
    assert api.status_code == 403
    assert api.get_json()["error"] == "password_change_required"
    assert password_page.status_code == 200


def test_admin_changes_password_and_can_enter_system(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()
    _login(client)

    missing_csrf = client.post(
        "/api/account/password",
        json={
            "current_password": "admin",
            "new_password": "Admin1234",
        },
    )
    wrong_current = _change_password(
        client, "wrong-current", "Admin1234"
    )
    changed = _change_password(client, "admin", "Admin1234")

    assert missing_csrf.status_code == 403
    assert missing_csrf.get_json()["error"] == "invalid_csrf_token"
    assert wrong_current.status_code == 400
    assert changed.status_code == 200
    assert client.get("/").status_code == 200
    assert client.get("/api/status").status_code == 200
    session_data = client.get("/api/session").get_json()
    assert session_data["operator"] == "系统管理员"
    assert session_data["role"] == "super_admin"
    assert session_data["can_manage_accounts"] is True


def test_super_admin_deletes_experiment_and_workbench_shows_creator_account(
    tmp_path,
):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    session_data = admin.get("/api/session").get_json()
    created = admin.post(
        "/api/experiments", json=_experiment_request()
    ).get_json()

    workbench = admin.get("/api/workbench").get_json()
    listed = next(
        item for item in workbench["experiments"]
        if item["id"] == created["id"]
    )
    missing_csrf = admin.delete(
        f"/api/experiments/{created['id']}"
    )
    deleted = admin.delete(
        f"/api/experiments/{created['id']}",
        headers={"X-CSRF-Token": session_data["csrf_token"]},
    )

    assert session_data["can_delete_experiments"] is True
    assert listed["created_by_username"] == "admin"
    assert listed["created_by_display_name"] == "系统管理员"
    assert missing_csrf.status_code == 403
    assert missing_csrf.get_json()["error"] == "invalid_csrf_token"
    assert deleted.status_code == 200
    assert deleted.get_json()["batch_id"] == created["batch_id"]
    assert repo.experiments.get_experiment(created["id"]) is None


def test_supervisor_and_operator_cannot_delete_experiments(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    admin_session = admin.get("/api/session").get_json()
    created = admin.post(
        "/api/experiments", json=_experiment_request()
    ).get_json()

    for username, display_name, role in (
        ("lab-supervisor", "实验室主管甲", "supervisor"),
        ("lab-operator", "实验操作员甲", "operator"),
    ):
        response = admin.post(
            "/api/accounts",
            json={
                "username": username,
                "display_name": display_name,
                "password": "Initial123",
                "role": role,
            },
            headers={
                "X-CSRF-Token": admin_session["csrf_token"]
            },
        )
        assert response.status_code == 201
        user = app.test_client()
        _login(user, username, "Initial123")
        _change_password(user, "Initial123", "Changed123")
        user_session = user.get("/api/session").get_json()

        denied = user.delete(
            f"/api/experiments/{created['id']}",
            headers={
                "X-CSRF-Token": user_session["csrf_token"]
            },
        )

        assert user_session["can_delete_experiments"] is False
        assert denied.status_code == 403
        assert denied.get_json()["error"] == "forbidden"

    assert repo.experiments.get_experiment(created["id"]) is not None


def test_login_rejects_invalid_credentials_and_rate_limits(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()

    responses = [
        _login(client, "admin", "wrong-password")
        for _ in range(6)
    ]

    assert [response.status_code for response in responses[:5]] == [
        401, 401, 401, 401, 401
    ]
    assert responses[0].get_json()["error"] == "invalid_credentials"
    assert responses[5].status_code == 429
    assert responses[5].get_json()["error"] == "too_many_attempts"


def test_super_admin_creates_supervisor_and_operator_accounts(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()
    _login(client)
    _change_password(client, "admin", "Admin1234")
    csrf = client.get("/api/session").get_json()["csrf_token"]

    supervisor = client.post(
        "/api/accounts",
        json={
            "username": "lablead",
            "display_name": "实验室主管",
            "password": "Lead12345",
            "role": "supervisor",
        },
        headers={"X-CSRF-Token": csrf},
    )
    operator = client.post(
        "/api/accounts",
        json={
            "username": "operator01",
            "display_name": "操作员一号",
            "password": "Operator123",
            "role": "operator",
        },
        headers={"X-CSRF-Token": csrf},
    )
    duplicate = client.post(
        "/api/accounts",
        json={
            "username": "OPERATOR01",
            "display_name": "重复账号",
            "password": "Operator456",
            "role": "operator",
        },
        headers={"X-CSRF-Token": csrf},
    )

    assert supervisor.status_code == 201
    assert supervisor.get_json()["must_change_password"] is True
    assert "password_hash" not in supervisor.get_json()
    assert operator.status_code == 201
    assert duplicate.status_code == 409
    account_data = client.get("/api/accounts").get_json()
    users = account_data["users"]
    assert account_data["creatable_roles"] == [
        "operator", "supervisor"
    ]
    assert {user["role"] for user in users} == {
        "super_admin", "supervisor", "operator"
    }


def test_supervisor_can_manage_only_operators(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    admin_csrf = admin.get("/api/session").get_json()["csrf_token"]
    admin.post(
        "/api/accounts",
        json={
            "username": "labboss",
            "display_name": "李主管",
            "password": "Lead12345",
            "role": "supervisor",
        },
        headers={"X-CSRF-Token": admin_csrf},
    )

    supervisor = app.test_client()
    _login(supervisor, "labboss", "Lead12345")
    _change_password(supervisor, "Lead12345", "Lead67890")
    supervisor_csrf = supervisor.get("/api/session").get_json()["csrf_token"]
    create_operator = supervisor.post(
        "/api/accounts",
        json={
            "username": "worker01",
            "display_name": "王操作员",
            "password": "Worker123",
            "role": "operator",
        },
        headers={"X-CSRF-Token": supervisor_csrf},
    )
    create_supervisor = supervisor.post(
        "/api/accounts",
        json={
            "username": "otherlead",
            "display_name": "另一主管",
            "password": "Other1234",
            "role": "supervisor",
        },
        headers={"X-CSRF-Token": supervisor_csrf},
    )

    assert create_operator.status_code == 201
    assert create_supervisor.status_code == 403
    visible = supervisor.get("/api/accounts").get_json()["users"]
    assert {user["role"] for user in visible} == {
        "supervisor", "operator"
    }


def test_operator_cannot_open_account_management(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    csrf = admin.get("/api/session").get_json()["csrf_token"]
    admin.post(
        "/api/accounts",
        json={
            "username": "tablet01",
            "display_name": "平板操作员",
            "password": "Tablet123",
            "role": "operator",
        },
        headers={"X-CSRF-Token": csrf},
    )

    operator = app.test_client()
    _login(operator, "tablet01", "Tablet123")
    _change_password(operator, "Tablet123", "Tablet456")

    assert operator.get("/experiments").status_code == 200
    assert operator.get("/accounts").status_code == 403
    assert operator.get("/api/accounts").status_code == 403
    assert operator.get("/change-password").status_code == 200


def test_same_display_name_cannot_take_over_batch_or_spoof_run_tag(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    csrf = admin.get("/api/session").get_json()["csrf_token"]
    admin.post(
        "/api/accounts",
        json={
            "username": "same-name-worker",
            "display_name": "系统管理员",
            "password": "Worker123",
            "role": "operator",
        },
        headers={"X-CSRF-Token": csrf},
    )
    experiment = admin.post(
        "/api/experiments",
        json={
            "batch_id": "20260724-CEM-01",
            "membrane_system": "CEM",
            "recipe_no": "R-CEM-001",
            "recipe_version": "V1",
            "sop_code": "SOP-SOL-GEL-CEM-AEM-01",
            "sop_version": "V0.2",
            "target_viscosity_min_mpas": 2.5,
            "target_viscosity_max_mpas": 10,
            "reviewer": "",
            "spec_snapshot": {},
        },
    ).get_json()

    worker = app.test_client()
    _login(worker, "same-name-worker", "Worker123")
    _change_password(worker, "Worker123", "Worker456")
    denied = worker.post(
        f"/api/experiments/{experiment['id']}/steps/R201-01/start",
        json={
            "row_version": experiment["row_version"],
            "client_event_id": "same-name-takeover",
        },
    )

    run_id = repo.open_run(did, 1, 1000, {})
    repo.close_run(run_id, 2000, "completed", 1, "mL", 1, "mL", 0)
    tagged = worker.post(
        f"/api/runs/{run_id}/tag",
        json={"operator": "伪造姓名", "project_tag": "验证"},
    )
    run = repo.get_run(run_id)

    assert denied.status_code == 403
    assert tagged.status_code == 200
    assert run.operator == "系统管理员"
    assert run.operator_user_id == worker.get(
        "/api/session"
    ).get_json()["user_id"]


def test_admin_can_reset_and_disable_subordinate_account(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    admin = app.test_client()
    _login(admin)
    _change_password(admin, "admin", "Admin1234")
    csrf = admin.get("/api/session").get_json()["csrf_token"]
    created = admin.post(
        "/api/accounts",
        json={
            "username": "disabled01",
            "display_name": "待停用操作员",
            "password": "Enabled123",
            "role": "operator",
        },
        headers={"X-CSRF-Token": csrf},
    ).get_json()

    reset = admin.post(
        f"/api/accounts/{created['id']}/reset-password",
        json={"password": "Reset1234"},
        headers={"X-CSRF-Token": csrf},
    )
    disabled = admin.patch(
        f"/api/accounts/{created['id']}/status",
        json={"is_active": False},
        headers={"X-CSRF-Token": csrf},
    )
    user_client = app.test_client()
    rejected = _login(user_client, "disabled01", "Reset1234")

    assert reset.status_code == 200
    assert reset.get_json()["must_change_password"] is True
    assert disabled.status_code == 200
    assert disabled.get_json()["is_active"] is False
    assert rejected.status_code == 401


def test_logout_clears_database_account_session(tmp_path):
    app, repo, did = _authenticated_app(tmp_path)
    client = app.test_client()
    _login(client)

    response = client.post("/logout")

    assert response.status_code == 302
    assert response.headers["Location"] == "/login"
    assert client.get("/api/session").status_code == 401
